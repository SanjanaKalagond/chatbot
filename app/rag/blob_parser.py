import magic
import pdfplumber
import docx
import openpyxl
import pytesseract
from PIL import Image
import io


def extract_text_from_blob(file_path):
    mime = magic.from_file(file_path, mime=True)

    text = ""

    if "pdf" in mime:
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text
        except:
            pass

    elif "word" in mime or "officedocument.wordprocessingml" in mime:
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        except:
            pass

    elif "spreadsheet" in mime or "excel" in mime:
        try:
            wb = openpyxl.load_workbook(file_path)
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(c) for c in row if c]) + "\n"
        except:
            pass

    elif "image" in mime:
        try:
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img)
        except:
            pass

    return text