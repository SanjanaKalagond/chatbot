import os
import boto3
import pickle
import faiss
import numpy as np
import sys
import magic
import pdfplumber
import docx
import openpyxl
import pytesseract
from PIL import Image
from sentence_transformers import SentenceTransformer
from sqlalchemy import text
from app.database.postgres import engine

model = SentenceTransformer("all-MiniLM-L6-v2")

s3 = boto3.client("s3")
BUCKET = "sf-chatbot-data"
SAVE_DIR = "data/faiss_index"


def extract_text(file_path):

    try:
        mime = magic.from_file(file_path, mime=True)
    except Exception:
        return ""

    text = ""

    try:

        if "pdf" in mime:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text

        elif "word" in mime or "officedocument.wordprocessingml" in mime:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"

        elif "excel" in mime or "spreadsheet" in mime:
            wb = openpyxl.load_workbook(file_path)
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(c) for c in row if c])
                    text += row_text + "\n"

        elif "text" in mime or "plain" in mime:
            with open(file_path, "r", errors="ignore") as f:
                text = f.read()

        elif "image" in mime:
            try:
                img = Image.open(file_path)
                text = pytesseract.image_to_string(img)
            except:
                pass

    except Exception:
        pass

    return text


def build_index():

    print("Fetching document paths from RDS...")
    sys.stdout.flush()

    with engine.connect() as conn:
        records = conn.execute(text("SELECT s3_path, id FROM documents")).fetchall()

    if not records:
        print("No documents found in database.")
        return

    all_chunks = []
    all_metadata = []

    for s3_path, doc_id in records:

        key = s3_path.replace(f"s3://{BUCKET}/", "")

        extension = key.split(".")[-1]
        local_path = f"/tmp/{doc_id}.{extension}"

        try:

            print(f"Downloading: {key}")
            sys.stdout.flush()

            s3.download_file(BUCKET, key, local_path)

            if not os.path.exists(local_path) or os.path.getsize(local_path) == 0:
                os.remove(local_path)
                continue

            full_text = extract_text(local_path)

            if not full_text or not full_text.strip():
                os.remove(local_path)
                continue

            start = 0
            size = 500
            overlap = 100

            while start < len(full_text):

                end = start + size
                chunk = full_text[start:end]

                if chunk.strip():
                    all_chunks.append(chunk)
                    all_metadata.append(
                        {
                            "doc_id": str(doc_id),
                            "text": chunk
                        }
                    )

                start = end - overlap

            os.remove(local_path)

        except Exception as e:

            print(f"Skipped {key} : {str(e)}")

            if os.path.exists(local_path):
                os.remove(local_path)

        sys.stdout.flush()

    if not all_chunks:
        print("No valid text extracted from documents.")
        return

    print(f"Generating embeddings for {len(all_chunks)} chunks...")
    sys.stdout.flush()

    embeddings = model.encode(all_chunks)
    embeddings = np.array(embeddings).astype("float32")

    dimension = embeddings.shape[1]

    index = faiss.IndexFlatL2(dimension)
    index.add(embeddings)

    os.makedirs(SAVE_DIR, exist_ok=True)

    index_path = f"{SAVE_DIR}/index.faiss"
    meta_path = f"{SAVE_DIR}/meta.pkl"

    faiss.write_index(index, index_path)

    with open(meta_path, "wb") as f:
        pickle.dump(all_metadata, f)

    print("FAISS index created successfully")
    print(f"Index saved at: {index_path}")
    print(f"Metadata saved at: {meta_path}")
    print(f"Total chunks indexed: {len(all_chunks)}")

    sys.stdout.flush()


if __name__ == "__main__":
    build_index()